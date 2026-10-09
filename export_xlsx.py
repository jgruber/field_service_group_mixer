"""Build the printable workbook from a congregation database.

Two sheets:

  Field Service Groups     the groups laid out three across, each block giving
                           its overseer and assistant, when and where it meets,
                           and the families assigned to it
  Meetings for Field Service   every one of those meetings, plus the
                           congregation's own, laid out as a week

Everything here reads; nothing is written back to the database.
"""

import re
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLUMNS   = 3          # group blocks across the first sheet
COL_WIDTH = 34
DAYS      = ['Sunday', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']

HEADER_FILL = PatternFill('solid', fgColor='4338CA')
DAY_FILL    = PatternFill('solid', fgColor='EEF2FF')
WHITE_BOLD  = Font(bold=True, color='FFFFFF', size=12)
GROUP_FONT  = Font(bold=True, size=11)
LABEL_FONT  = Font(size=10)
MUTED_FONT  = Font(size=9, color='666666')
TIME_FONT   = Font(bold=True, size=10)
THIN        = Side(style='thin', color='BFBFBF')


# ── reading ────────────────────────────────────────────────────────────────

def _table_exists(cur, name):
    cur.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (name,))
    return cur.fetchone() is not None


def read_model(db_path):
    """Everything the workbook needs, shaped the way the app shows it."""
    con = sqlite3.connect(db_path)
    con.row_factory = sqlite3.Row
    cur = con.cursor()

    meetings = {}
    if _table_exists(cur, 'fsg_meetings'):
        for r in cur.execute('SELECT group_id, day, time, location FROM fsg_meetings'):
            meetings[r['group_id']] = dict(day=r['day'] or '', time=r['time'] or '',
                                           location=r['location'] or '')

    congregation_meetings = []
    if _table_exists(cur, 'congregation_meetings'):
        for r in cur.execute('SELECT day, time, location FROM congregation_meetings ORDER BY position'):
            congregation_meetings.append(dict(day=r['day'] or '', time=r['time'] or '',
                                              location=r['location'] or ''))

    # Only families the app itself shows: not moved, and with someone still in them.
    cur.execute("""
        SELECT f.id, f.name, f.family_head, f.field_service_group_id
          FROM families f
         WHERE f.moved = 0
           AND EXISTS (SELECT 1 FROM persons p
                        WHERE p.family_id = f.id AND p.moved = 0
                          AND (p.removed IS NULL OR p.removed = 0))
    """)
    by_group = {}
    for r in cur.fetchall():
        by_group.setdefault(r['field_service_group_id'], []).append(
            dict(name=r['name'] or '(Unknown)', head=r['family_head'] or ''))
    for fams in by_group.values():
        fams.sort(key=lambda f: f['name'].lower())

    groups = []
    for r in cur.execute("""SELECT id, name, overseer, assistant
                              FROM field_service_groups ORDER BY name COLLATE NOCASE"""):
        groups.append(dict(id=r['id'], name=r['name'] or '(Unnamed)',
                           overseer=r['overseer'] or '', assistant=r['assistant'] or '',
                           families=by_group.get(r['id'], []),
                           **meetings.get(r['id'], dict(day='', time='', location=''))))

    known = {g['id'] for g in groups}
    stray = [f for gid, fams in by_group.items() if gid not in known for f in fams]
    if stray:
        groups.append(dict(id=None, name='Unassigned', overseer='', assistant='',
                           families=sorted(stray, key=lambda f: f['name'].lower()),
                           day='', time='', location=''))

    # Addresses, so a location that names a household can be split over lines.
    addresses = {}
    for r in cur.execute("""SELECT DISTINCT full_address, address, city, state, postal_code
                              FROM persons
                             WHERE full_address IS NOT NULL AND full_address <> ''"""):
        addresses[r['full_address'].strip().lower()] = [
            (r['address'] or '').strip(),
            ', '.join(x for x in [(r['city'] or '').strip(),
                                  ' '.join(y for y in [(r['state'] or '').strip(),
                                                       (r['postal_code'] or '').strip()] if y)] if x),
        ]

    name = 'Congregation'
    row = cur.execute('SELECT name FROM congregations LIMIT 1').fetchone()
    if row and row['name']:
        name = row['name']

    con.close()
    return dict(congregation=name, groups=groups,
                congregation_meetings=congregation_meetings, addresses=addresses)


# ── formatting helpers ─────────────────────────────────────────────────────

def format_time(value):
    """'14:30' as it would be written on a notice board."""
    m = re.match(r'^(\d{1,2}):(\d{2})', value or '')
    if not m:
        return value or ''
    hour = int(m.group(1))
    return '%d:%s %s' % (hour % 12 or 12, m.group(2), 'AM' if hour < 12 else 'PM')


def format_when(day, time):
    return ' '.join(x for x in [day, format_time(time)] if x)


STATE_ZIP = re.compile(r'^[A-Za-z]{2}\.?\s+\d{5}(?:-\d{4})?$')
ZIP_ONLY  = re.compile(r'^\d{5}(?:-\d{4})?$')


def normalize_address(text):
    """The same spelling rule the page applies: no comma after the street."""
    parts = [p.strip() for p in str(text or '').split(',') if p.strip()]
    if len(parts) < 2:
        return str(text or '').strip()
    if not (STATE_ZIP.match(parts[-1]) or ZIP_ONLY.match(parts[-1])):
        return ', '.join(parts)
    return '%s, %s' % (' '.join(parts[:-1]), parts[-1])


def format_location(location, addresses):
    """A location over as many lines as it naturally has.

    An address that names a household in the database is split the way the
    database holds it. Anything else is split on its commas — except that a
    trailing "TX 75025" rejoins the city before it, so a typed address still
    reads as street / city, state zip rather than breaking after the city.
    """
    text = (location or '').strip()
    if not text:
        return []
    # A typed address may carry a comma the database does not use; try both
    # spellings before falling back to splitting the text as written.
    known = addresses.get(text.lower()) or addresses.get(normalize_address(text).lower())
    if known:
        return [line for line in known if line]

    lines = []
    for part in (p.strip() for p in text.split(',')):
        if not part:
            continue
        if lines and (STATE_ZIP.match(part) or ZIP_ONLY.match(part)):
            lines[-1] = '%s, %s' % (lines[-1], part)
        else:
            lines.append(part)
    return lines


def format_family(fam):
    """'Gruber (John)' — the family name, then the head's first name."""
    first = (fam['head'] or '').strip().split(' ')[0]
    return '%s (%s)' % (fam['name'], first) if first else fam['name']


def group_block(group, addresses):
    """One group as two stacks of lines: who and when above, families below.

    They are kept apart so that a band of three groups can have its family
    lists start on the same row even when one of them has a longer address.
    """
    head = [(group['name'], 'group'),
            (group['overseer'] or '—', 'person'),
            (group['assistant'] or '—', 'person'),
            ('', 'blank'),
            (format_when(group['day'], group['time']) or 'Meeting not set', 'when')]
    head += [(part, 'where') for part in format_location(group['location'], addresses)]

    if group['families']:
        families = [(format_family(fam), 'family') for fam in group['families']]
    else:
        families = [('No families', 'where')]
    return head, families


# ── sheet one ──────────────────────────────────────────────────────────────

def write_groups_sheet(ws, model):
    ws.title = 'Field Service Groups'
    for i in range(COLUMNS):
        ws.column_dimensions[get_column_letter(i + 1)].width = COL_WIDTH

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COLUMNS)
    title = ws.cell(row=1, column=1, value='%s — Field Service Groups' % model['congregation'])
    title.font = WHITE_BOLD
    title.fill = HEADER_FILL
    title.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 22

    styles = {
        'group':  (GROUP_FONT, None),
        'person': (LABEL_FONT, None),
        'when':   (TIME_FONT, None),
        'where':  (MUTED_FONT, None),
        'family': (LABEL_FONT, None),
        'blank':  (LABEL_FONT, None),
    }

    row = 3
    for start in range(0, len(model['groups']), COLUMNS):
        blocks = [group_block(g, model['addresses']) for g in model['groups'][start:start + COLUMNS]]
        # Pad every header to the tallest in the band, then a blank line, so the
        # family lists below all begin on the same row.
        head_height = max(len(head) for head, _ in blocks) + 1
        padded = [head + [('', 'blank')] * (head_height - len(head)) + families
                  for head, families in blocks]
        height = max(len(b) for b in padded)
        for col, block in enumerate(padded, start=1):
            for offset in range(height):
                cell = ws.cell(row=row + offset, column=col)
                if offset < len(block):
                    text, kind = block[offset]
                    cell.value = text
                    cell.font = styles[kind][0]
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.border = Border(
                    left=THIN,
                    right=THIN,
                    top=THIN if offset == 0 else None,
                    bottom=THIN if offset == height - 1 else None,
                )
        row += height + 1          # a spacer row between bands of groups

    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:1'


# ── sheet two ──────────────────────────────────────────────────────────────

def collect_week(model):
    """Every meeting, grouped by day, each as (time, label, location)."""
    week = {day: [] for day in DAYS}
    unscheduled = []

    def add(day, time, label, location):
        entry = (time or '', label, location or '')
        (week[day] if day in week else unscheduled).append(entry)

    for m in model['congregation_meetings']:
        add(m['day'], m['time'], 'Congregation', m['location'])
    for g in model['groups']:
        if g['day'] or g['time'] or g['location']:
            add(g['day'], g['time'], g['name'], g['location'])

    for entries in week.values():
        entries.sort(key=lambda e: (e[0] == '', e[0]))
    return week, unscheduled


def write_week_sheet(ws, model):
    ws.title = 'Meetings for Field Service'
    week, unscheduled = collect_week(model)

    for i in range(len(DAYS)):
        ws.column_dimensions[get_column_letter(i + 1)].width = 26

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(DAYS))
    title = ws.cell(row=1, column=1, value='%s — Meetings for Field Service' % model['congregation'])
    title.font = WHITE_BOLD
    title.fill = HEADER_FILL
    title.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 22

    for i, day in enumerate(DAYS, start=1):
        cell = ws.cell(row=2, column=i, value=day)
        cell.font = Font(bold=True, size=11)
        cell.fill = DAY_FILL
        cell.alignment = Alignment(horizontal='center')
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    columns = []
    for day in DAYS:
        lines = []
        for time, label, location in week[day]:
            lines.append(('%s  %s' % (format_time(time), label) if time else label, 'when'))
            for part in format_location(location, model['addresses']):
                lines.append((part, 'where'))
            lines.append(('', 'blank'))
        columns.append(lines[:-1] if lines else [])

    height = max([len(c) for c in columns] + [1])
    for col, lines in enumerate(columns, start=1):
        for offset in range(height):
            cell = ws.cell(row=3 + offset, column=col)
            if offset < len(lines):
                text, kind = lines[offset]
                cell.value = text
                cell.font = TIME_FONT if kind == 'when' else MUTED_FONT
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(
                left=THIN, right=THIN,
                top=THIN if offset == 0 else None,
                bottom=THIN if offset == height - 1 else None,
            )

    row = 3 + height + 1
    if unscheduled:
        ws.cell(row=row, column=1, value='No day set').font = Font(bold=True, size=11)
        row += 1
        for time, label, location in unscheduled:
            where = ', '.join(format_location(location, model['addresses']))
            ws.cell(row=row, column=1,
                    value=' '.join(x for x in [format_time(time), label] if x)).font = TIME_FONT
            if where:
                ws.cell(row=row, column=2, value=where).font = MUTED_FONT
            row += 1

    ws.freeze_panes = 'A3'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_workbook(db_path):
    model = read_model(db_path)
    wb = Workbook()
    write_groups_sheet(wb.active, model)
    write_week_sheet(wb.create_sheet(), model)
    return wb
